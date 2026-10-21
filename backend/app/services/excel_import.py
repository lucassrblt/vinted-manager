"""Import unique de l'Excel « Suivi revente » : lecture -> aperçu corrigeable -> écriture.

`lire_classeur` ne touche jamais à la base. `ecrire_import` écrit tout dans une seule transaction.
"""

import io
import re
from collections import Counter, defaultdict
from collections.abc import Mapping
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Article, Batch, HistoriqueStatut, PisteSourcing
from app.models.enums import StatutArticle
from app.services import calculs, marques

ONGLET_SUIVI = "Suivi revente"
ONGLET_BATCH_2 = "Batch 2"
ONGLET_SOURCING_TESTS = "Sourcing tests"
ONGLET_SOURCING_RECHERCHE = "Sourcing - Recherche"

STATUTS_EXCEL: dict[str, StatutArticle] = {
    "vendu": StatutArticle.VENDU,
    "en ligne": StatutArticle.EN_LIGNE,
    "recu / a publier": StatutArticle.RECU_A_PUBLIER,
    "commande / non recu": StatutArticle.COMMANDE,
}

# En-tête Excel (normalisé) -> champ de LigneImport.
COLONNES_SUIVI: dict[str, str] = {
    "nom de l'article": "nom",
    "taille": "taille",
    "prix d'achat": "prix_achat_net",
    "prix de vente": "prix_vente",
    "vendu (oui/non)": "vendu",
    "litige (oui/non)": "litige",
    "batch": "batch",
    "source": "source",
    "prix achat brut": "prix_achat_brut",
    "remise": "remise",
    "statut": "statut_excel",
    "date d'achat": "date_achat",
    "date achat": "date_achat",
}

Niveau = Literal["erreur", "avertissement"]


class Anomalie(BaseModel):
    champ: str
    niveau: Niveau
    message: str


class LigneImport(BaseModel):
    ligne: int
    nom: str = ""
    taille: str | None = None
    marque: str | None = None
    categorie: str | None = None
    couleur: str | None = None
    prix_achat_brut: Decimal | None = None
    remise: Decimal = Decimal("0")
    prix_achat_net: Decimal | None = None
    prix_vente: Decimal | None = None
    vendu: bool = False
    litige: bool = False
    batch: str | None = None
    source: str | None = None
    date_achat: date | None = None
    statut_excel: str | None = None
    statut: StatutArticle | None = None
    anomalies: list[Anomalie] = Field(default_factory=list)


class PisteImport(BaseModel):
    origine: Literal["tests", "recherche"]
    marque: str | None = None
    produit: str
    mots_cles: str | None = None
    categorie: str | None = None
    achat_max: Decimal | None = None
    revente_basse: Decimal | None = None
    revente_haute: Decimal | None = None
    priorite: str | None = None
    quantite_test: str | None = None
    commentaire: str | None = None


class ApercuImport(BaseModel):
    lignes: list[LigneImport]
    pistes: list[PisteImport]
    onglets_lus: list[str]
    onglets_absents: list[str]


class ErreurImport(ValueError):
    pass


# --- Conversions de cellules ----------------------------------------------------------------


def _texte(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _montant(v: Any) -> Decimal | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, int | float):
        return Decimal(str(round(float(v), 2))).quantize(Decimal("0.01"))
    s = re.sub(r"[€\s ]", "", str(v)).replace(",", ".")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _booleen(v: Any) -> bool:
    return marques.normaliser(str(v or "")).strip() in {"oui", "o", "yes", "x", "true", "1", "vrai"}


def _date(v: Any) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.match(r"^\s*(\d{1,2})/(\d{1,2})/(\d{2,4})\s*$", str(v))
    if not m:
        return None
    j, mo, a = (int(x) for x in m.groups())
    try:
        return date(a + 2000 if a < 100 else a, mo, j)
    except ValueError:
        return None


def _cellule(row: tuple[Any, ...], index: Mapping[str, int | None], champ: str) -> Any:
    i = index.get(champ)
    return row[i] if i is not None and i < len(row) else None


def _entete(v: Any) -> str:
    return marques.normaliser(str(v or "")).strip().replace("’", "'")


# --- Lecture ---------------------------------------------------------------------------------


def _lire_suivi(ws: Any) -> list[LigneImport]:
    lignes_brutes = list(ws.iter_rows(values_only=True))
    if not lignes_brutes:
        raise ErreurImport(f"L'onglet « {ONGLET_SUIVI} » est vide.")
    entetes = [_entete(h) for h in lignes_brutes[0]]
    index: dict[str, int] = {}
    for i, h in enumerate(entetes):
        champ = COLONNES_SUIVI.get(h)
        if champ and champ not in index:
            index[champ] = i
    manquantes = [c for c in ("nom", "prix_achat_net", "statut_excel") if c not in index]
    if manquantes:
        raise ErreurImport(
            f"Colonnes introuvables dans « {ONGLET_SUIVI} » : {', '.join(manquantes)}. "
            "Vérifie que la première ligne contient bien les en-têtes."
        )
    # La date d'achat est dans une colonne sans en-tête, juste après « Source ».
    if "date_achat" not in index and "source" in index:
        candidat = index["source"] + 1
        if candidat < len(entetes) and entetes[candidat] == "":
            index["date_achat"] = candidat

    resultat: list[LigneImport] = []
    for num, row in enumerate(lignes_brutes[1:], start=2):

        def cell(champ: str, row: tuple[Any, ...] = row) -> Any:
            return _cellule(row, index, champ)

        nom = _texte(cell("nom"))
        if not nom:
            continue
        prix_net = _montant(cell("prix_achat_net"))
        brut = _montant(cell("prix_achat_brut"))
        remise = _montant(cell("remise")) or Decimal("0")
        if brut is None and prix_net is not None:
            brut = prix_net + remise
        statut_excel = _texte(cell("statut_excel"))
        resultat.append(
            LigneImport(
                ligne=num,
                nom=nom,
                taille=_texte(cell("taille")),
                marque=marques.detecter_marque(nom),
                categorie=marques.detecter_categorie(nom),
                couleur=marques.detecter_couleur(nom),
                prix_achat_brut=brut,
                remise=remise,
                prix_achat_net=prix_net,
                prix_vente=_montant(cell("prix_vente")),
                vendu=_booleen(cell("vendu")),
                litige=_booleen(cell("litige")),
                batch=_texte(cell("batch")),
                source=_texte(cell("source")),
                date_achat=_date(cell("date_achat")),
                statut_excel=statut_excel,
                statut=STATUTS_EXCEL.get(marques.normaliser(statut_excel or "").strip()),
            )
        )
    return resultat


def _prix_cibles_batch_2(ws: Any) -> dict[tuple[str, str, str], Decimal]:
    """Onglet « Batch 2 » : (produit, taille, couleur) -> prix de vente cible."""
    cibles: dict[tuple[str, str, str], Decimal] = {}
    index: dict[str, int | None] = {}
    for row in ws.iter_rows(values_only=True):
        if not index:
            entetes = [_entete(v) for v in row]
            if "produit" in entetes and "prix vente cible" in entetes:
                index = {h: i for i, h in enumerate(entetes) if h}
            continue
        produit = _texte(_cellule(row, index, "produit"))
        prix = _montant(_cellule(row, index, "prix vente cible"))
        if produit and prix is not None:
            cle = (
                marques.normaliser(produit),
                marques.normaliser(_texte(_cellule(row, index, "taille")) or ""),
                marques.normaliser(_texte(_cellule(row, index, "couleur")) or ""),
            )
            cibles.setdefault(cle, prix)
    return cibles


def _lire_pistes(ws: Any, origine: Literal["tests", "recherche"]) -> list[PisteImport]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    entetes = [_entete(h) for h in rows[0]]

    def idx(*noms: str) -> int | None:
        for n in noms:
            if n in entetes:
                return entetes.index(n)
        return None

    if origine == "tests":
        cols = {
            "produit": idx("article"),
            "achat_max": idx("prix hacoo", "prix achat"),
            "revente_basse": idx("vente basse"),
            "revente_haute": idx("vente haute"),
            "priorite": idx("priorite"),
            "quantite_test": idx("qte test"),
            "commentaire": idx("commentaire"),
        }
    else:
        cols = {
            "marque": idx("marque"),
            "produit": idx("produit complet a chercher", "produit"),
            "mots_cles": idx("mots-cles a tester", "mots-cles"),
            "categorie": idx("categorie"),
            "achat_max": idx("achat max cible"),
            "revente_basse": idx("revente cible basse"),
            "revente_haute": idx("revente cible haute"),
            "priorite": idx("priorite"),
            "quantite_test": idx("qte test"),
            "commentaire": idx("note"),
        }

    pistes: list[PisteImport] = []
    for row in rows[1:]:

        def v(champ: str, row: tuple[Any, ...] = row) -> Any:
            return _cellule(row, cols, champ)

        produit = _texte(v("produit"))
        if not produit:
            continue
        qte = v("quantite_test")
        if isinstance(qte, float) and qte.is_integer():
            qte = int(qte)
        pistes.append(
            PisteImport(
                origine=origine,
                marque=_texte(v("marque")) or marques.detecter_marque(produit),
                produit=produit,
                mots_cles=_texte(v("mots_cles")),
                categorie=_texte(v("categorie")) or marques.detecter_categorie(produit),
                achat_max=_montant(v("achat_max")),
                revente_basse=_montant(v("revente_basse")),
                revente_haute=_montant(v("revente_haute")),
                priorite=_texte(v("priorite")),
                quantite_test=_texte(qte),
                commentaire=_texte(v("commentaire")),
            )
        )
    return pistes


def verifier_ligne(ligne: LigneImport) -> LigneImport:
    """(Re)calcule les anomalies d'une ligne : appelée à la lecture et à la validation."""
    a: list[Anomalie] = []
    if not ligne.nom.strip():
        a.append(Anomalie(champ="nom", niveau="erreur", message="Nom manquant"))
    if ligne.prix_achat_net is None:
        a.append(Anomalie(champ="prix_achat_net", niveau="erreur", message="Prix d'achat manquant"))
    if ligne.statut is None:
        a.append(
            Anomalie(
                champ="statut",
                niveau="erreur",
                message=f"Statut inconnu : « {ligne.statut_excel or 'vide'} »",
            )
        )
    if ligne.statut == StatutArticle.VENDU and ligne.prix_vente is None:
        a.append(Anomalie(champ="prix_vente", niveau="erreur", message="Article vendu sans prix de vente"))
    statut_origine = STATUTS_EXCEL.get(marques.normaliser(ligne.statut_excel or "").strip())
    if ligne.vendu and statut_origine not in (None, StatutArticle.VENDU):
        a.append(
            Anomalie(
                champ="statut",
                niveau="avertissement",
                message=f"« Vendu = oui » mais statut Excel « {ligne.statut_excel} » : à vérifier",
            )
        )
    if not ligne.taille:
        a.append(Anomalie(champ="taille", niveau="avertissement", message="Taille vide"))
    if not ligne.marque:
        a.append(Anomalie(champ="marque", niveau="avertissement", message="Marque non détectée"))
    if not ligne.batch:
        a.append(Anomalie(champ="batch", niveau="avertissement", message="Aucun batch"))
    ligne.anomalies = a
    return ligne


def lire_classeur(contenu: bytes) -> ApercuImport:
    try:
        wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lève des exceptions variées sur un fichier invalide
        raise ErreurImport("Le fichier n'est pas un classeur Excel (.xlsx) lisible.") from exc

    if ONGLET_SUIVI not in wb.sheetnames:
        raise ErreurImport(f"Onglet « {ONGLET_SUIVI} » introuvable dans le fichier.")
    lus, absents = [ONGLET_SUIVI], []
    lignes = _lire_suivi(wb[ONGLET_SUIVI])

    if ONGLET_BATCH_2 in wb.sheetnames:
        lus.append(ONGLET_BATCH_2)
        cibles = _prix_cibles_batch_2(wb[ONGLET_BATCH_2])
        for ligne in lignes:
            if ligne.batch == "Batch 2" and ligne.prix_vente is None and " - " in ligne.nom:
                produit, couleur = ligne.nom.rsplit(" - ", 1)
                cle = (
                    marques.normaliser(produit),
                    marques.normaliser(ligne.taille or ""),
                    marques.normaliser(couleur),
                )
                ligne.prix_vente = cibles.get(cle)

    for ligne in lignes:
        # Incohérence « Vendu = oui » / statut : la colonne « Vendu » fait foi dans les totaux
        # de l'Excel, on propose donc « vendu » par défaut (modifiable dans l'aperçu).
        if ligne.vendu and ligne.prix_vente is not None:
            ligne.statut = StatutArticle.VENDU
        verifier_ligne(ligne)

    pistes: list[PisteImport] = []
    for onglet, origine in ((ONGLET_SOURCING_TESTS, "tests"), (ONGLET_SOURCING_RECHERCHE, "recherche")):
        if onglet in wb.sheetnames:
            lus.append(onglet)
            pistes.extend(_lire_pistes(wb[onglet], origine))  # type: ignore[arg-type]
        else:
            absents.append(onglet)
    wb.close()
    return ApercuImport(lignes=lignes, pistes=pistes, onglets_lus=lus, onglets_absents=absents)


# --- Écriture --------------------------------------------------------------------------------


def base_deja_importee(db: Session) -> bool:
    return (db.scalar(select(func.count(Article.id))) or 0) > 0


def ecrire_import(db: Session, lignes: list[LigneImport], pistes: list[PisteImport]) -> dict[str, int]:
    lignes = [verifier_ligne(lg) for lg in lignes]
    bloquantes = [lg for lg in lignes if any(a.niveau == "erreur" for a in lg.anomalies)]
    if bloquantes:
        numeros = ", ".join(str(lg.ligne) for lg in bloquantes[:10])
        raise ErreurImport(f"{len(bloquantes)} ligne(s) contiennent encore des erreurs (lignes {numeros}).")
    if base_deja_importee(db):
        raise ErreurImport("Des articles existent déjà : l'import Excel ne peut être fait qu'une fois.")

    par_batch: dict[str, list[LigneImport]] = defaultdict(list)
    for lg in lignes:
        if lg.batch:
            par_batch[lg.batch].append(lg)

    batches: dict[str, Batch] = {}
    for nom, membres in par_batch.items():
        sources = Counter(m.source for m in membres if m.source)
        dates = [m.date_achat for m in membres if m.date_achat]
        batch = Batch(
            nom=nom,
            fournisseur=sources.most_common(1)[0][0] if sources else None,
            date_commande=min(dates) if dates else None,
            total_paye=sum((m.prix_achat_net or Decimal("0") for m in membres), Decimal("0")),
            remise_globale=sum((m.remise for m in membres), Decimal("0")),
        )
        db.add(batch)
        batches[nom] = batch

    for lg in lignes:
        assert lg.statut is not None
        vendu = lg.statut == StatutArticle.VENDU
        article = Article(
            nom=lg.nom.strip(),
            marque=lg.marque,
            categorie=lg.categorie,
            taille=lg.taille,
            couleur=lg.couleur,
            prix_achat_brut=lg.prix_achat_brut,
            remise=lg.remise,
            prix_achat_net=lg.prix_achat_net
            if lg.prix_achat_net is not None
            else calculs.prix_achat_net(lg.prix_achat_brut, lg.remise),
            prix_vente_cible=lg.prix_vente,
            prix_vente_reel=lg.prix_vente if vendu else None,
            statut=lg.statut,
            litige=lg.litige,
            batch=batches.get(lg.batch) if lg.batch else None,
            source=lg.source,
            date_achat=lg.date_achat,
            notes=f"Importé depuis l'Excel (ligne {lg.ligne})",
        )
        db.add(article)
        db.add(
            HistoriqueStatut(
                article=article,
                ancien_statut=None,
                nouveau_statut=lg.statut,
                commentaire="Import Excel",
            )
        )

    for p in pistes:
        db.add(PisteSourcing(**p.model_dump()))

    db.commit()
    return {"articles": len(lignes), "batches": len(batches), "pistes": len(pistes)}
