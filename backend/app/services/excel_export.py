"""Export Excel avec la même structure que l'Excel d'origine (sauvegarde / comptable)."""

import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Article, Batch, Parametres, PisteSourcing
from app.models.enums import StatutArticle
from app.services import calculs, stats

STATUTS_LIBELLES = {
    StatutArticle.COMMANDE: "Commandé / non reçu",
    StatutArticle.RECU_A_PUBLIER: "Reçu / à publier",
    StatutArticle.EN_LIGNE: "En ligne",
    StatutArticle.VENDU: "Vendu",
    StatutArticle.ENVOYE: "Envoyé",
    StatutArticle.LIVRE: "Livré",
    StatutArticle.LITIGE: "Litige",
    StatutArticle.RETOURNE: "Retourné",
}

ENTETES_SUIVI = [
    "Nom de l'article",
    "Taille",
    "Prix d'achat",
    "Prix de vente",
    "Vendu (oui/non)",
    "Litige (oui/non)",
    "Marge nette",
    "Marge %",
    "Coefficient",
    "Objectif 2-4x",
    "Batch",
    "Source",
    "Date d'achat",
    "Prix achat brut",
    "Remise",
    "Statut",
]
GRAS = Font(bold=True)
FOND_ENTETE = PatternFill("solid", fgColor="E4E4E7")
EUROS = '#,##0.00 "€"'
POURCENT = "0.0%"


def _entete(ws: Worksheet, valeurs: list[str]) -> None:
    ws.append(valeurs)
    for cell in ws[ws.max_row]:
        cell.font = GRAS
        cell.fill = FOND_ENTETE
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _largeurs(ws: Worksheet, largeurs: list[int]) -> None:
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur


def generer_export(db: Session, params: Parametres) -> bytes:
    articles = list(db.scalars(select(Article).order_by(Article.batch_id.nulls_last(), Article.id)))
    batches = {b.id: b for b in db.scalars(select(Batch))}
    wb = Workbook()

    # Onglet suivi : mêmes colonnes que l'Excel d'origine
    ws = wb.active
    assert ws is not None
    ws.title = "Suivi revente"
    _entete(ws, ENTETES_SUIVI)
    for a in articles:
        pv = stats.prix_vente(a)
        vendu = stats.est_vendu(a)
        # Comme dans l'Excel : marge et coefficient seulement pour les articles vendus.
        marge = calculs.marge_nette(pv, a.prix_achat_net) if vendu else None
        coeff = calculs.coefficient(pv, a.prix_achat_net) if vendu else None
        ind = calculs.indicateur(coeff, params.coefficient_min, params.coefficient_max)
        ws.append(
            [
                a.nom,
                a.taille,
                a.prix_achat_net,
                pv,
                "oui" if vendu else "non",
                "oui" if a.litige else "non",
                marge,
                calculs.marge_pct(pv, a.prix_achat_net) if vendu else None,
                coeff,
                ind.value if ind else "En stock",
                batches[a.batch_id].nom if a.batch_id in batches else None,
                a.source,
                a.date_achat,
                a.prix_achat_brut,
                a.remise,
                STATUTS_LIBELLES.get(StatutArticle(a.statut)),
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for idx in (2, 3, 6, 13, 14):
            row[idx].number_format = EUROS
        row[7].number_format = POURCENT
        row[8].number_format = "0.00"
        row[12].number_format = "DD/MM/YYYY"
    _largeurs(ws, [48, 8, 12, 12, 10, 10, 12, 10, 11, 16, 16, 22, 12, 13, 10, 20])

    # Onglet synthèse
    k = stats.kpis(articles, params, None, None)
    ws2 = wb.create_sheet("Synthèse")
    _entete(ws2, ["SYNTHÈSE GLOBALE", "Valeur"])
    lignes_synthese: list[tuple[str, object, str | None]] = [
        ("Total investi (vendu + stock)", k.total_investi, EUROS),
        ("Valeur du stock non vendu", k.valeur_stock, EUROS),
        ("Chiffre d'affaires encaissé", k.ca_encaisse, EUROS),
        ("Marge réalisée (articles vendus)", k.marge_realisee, EUROS),
        ("Marge actuelle (stock inclus)", k.marge_globale, EUROS),
        ("Taux fiscal estimé", k.taux_cotisations, POURCENT),
        ("Cotisations estimées sur CA", k.cotisations_estimees, EUROS),
        ("Marge après fisc (stock inclus)", k.marge_globale_apres_cotisations, EUROS),
        ("Coefficient actuel CA / investi", k.coefficient_global, "0.00"),
        ("Nombre d'articles total", k.nb_total, None),
        ("Articles vendus", k.nb_vendus, None),
        ("Articles en stock", k.nb_en_stock, None),
        ("Sell-through (vendus / publiés)", k.sell_through, POURCENT),
        ("Articles reçus", k.nb_recus, None),
        ("Articles publiés", k.nb_publies, None),
        ("Articles à publier", k.nb_a_publier, None),
    ]
    for libelle, valeur, fmt in lignes_synthese:
        ws2.append([libelle, valeur])
        if fmt:
            ws2.cell(ws2.max_row, 2).number_format = fmt
    _largeurs(ws2, [38, 16])

    # Onglet performance par batch
    ws3 = wb.create_sheet("Performance par batch")
    _entete(
        ws3,
        [
            "Batch",
            "Articles",
            "Publiés",
            "Vendus",
            "Sell-through publié",
            "CA vendu",
            "Marge réalisée",
            "Coeff. ventes",
            "Délai moyen (jours)",
            "Stock restant",
        ],
    )
    par_batch: dict[int, list[Article]] = defaultdict(list)
    for a in articles:
        if a.batch_id:
            par_batch[a.batch_id].append(a)
    for b in sorted(batches.values(), key=lambda b: b.nom):
        s = stats.stats_batch(b, par_batch.get(b.id, []))
        delai = s["delai_moyen_heures"]
        ws3.append(
            [
                b.nom,
                s["nb_articles"],
                s["nb_publies"],
                s["nb_vendus"],
                s["sell_through"],
                s["ca"],
                s["marge_realisee"],
                s["coefficient_ventes"],
                round(delai / 24, 1) if isinstance(delai, float) else None,
                s["valeur_stock_restant"],
            ]
        )
        r = ws3.max_row
        ws3.cell(r, 5).number_format = POURCENT
        for c in (6, 7, 10):
            ws3.cell(r, c).number_format = EUROS
        ws3.cell(r, 8).number_format = "0.00"
    _largeurs(ws3, [20, 10, 10, 10, 18, 12, 14, 12, 18, 14])

    # Onglets sourcing
    pistes = list(db.scalars(select(PisteSourcing).order_by(PisteSourcing.id)))
    ws4 = wb.create_sheet("Sourcing")
    _entete(
        ws4,
        [
            "Origine",
            "Marque",
            "Produit",
            "Mots-clés",
            "Catégorie",
            "Achat max",
            "Revente basse",
            "Revente haute",
            "Priorité",
            "Qté test",
            "Commentaire",
        ],
    )
    for p in pistes:
        ws4.append(
            [
                p.origine,
                p.marque,
                p.produit,
                p.mots_cles,
                p.categorie,
                p.achat_max,
                p.revente_basse,
                p.revente_haute,
                p.priorite,
                p.quantite_test,
                p.commentaire,
            ]
        )
    _largeurs(ws4, [10, 16, 40, 40, 18, 11, 13, 13, 11, 9, 50])

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
