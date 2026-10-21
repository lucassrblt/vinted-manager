"""Import de l'Excel réel (tests/data/suivi-revente.xlsx, non versionné)."""

import io
from collections import Counter
from decimal import Decimal as D

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Article, Batch, HistoriqueStatut, PisteSourcing
from app.models.enums import StatutArticle
from app.services.excel_import import lire_classeur

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _apercu(client: TestClient, contenu: bytes) -> dict:
    r = client.post("/api/import/excel/preview", files={"fichier": ("suivi.xlsx", contenu, XLSX)})
    assert r.status_code == 200, r.text
    return r.json()


def test_lecture_sans_ecriture(client: TestClient, db: Session, excel_reel: bytes) -> None:
    apercu = _apercu(client, excel_reel)
    assert len(apercu["lignes"]) == 58
    assert db.scalar(select(func.count(Article.id))) == 0  # rien n'est écrit avant validation


def test_lecture_du_classeur(excel_reel: bytes) -> None:
    a = lire_classeur(excel_reel)
    statuts = Counter(lg.statut for lg in a.lignes)
    # 24 « Vendu » + 2 lignes « Vendu = oui » au statut incohérent, proposées en « vendu »
    assert statuts[StatutArticle.VENDU] == 26
    assert statuts[StatutArticle.COMMANDE] == 9
    assert len({lg.batch for lg in a.lignes}) == 5
    incoherentes = [lg for lg in a.lignes if any("Vendu = oui" in x.message for x in lg.anomalies)]
    assert {lg.nom for lg in incoherentes} == {
        "Patagonia colorblock / Snap-T style",
        "Carhartt demi-zip sweat",
    }
    # Date d'achat lue dans la colonne sans en-tête
    pull = next(lg for lg in a.lignes if lg.batch == "Batch 2")
    assert pull.date_achat is not None and pull.date_achat.strftime("%d/%m/%Y") == "02/10/2026"
    assert pull.prix_achat_brut == D("17.64") and pull.remise == D("0.97")
    # Batch 1 n'a pas de brut : brut = net, remise 0
    b1 = next(lg for lg in a.lignes if lg.batch == "Batch 1")
    assert b1.prix_achat_brut == b1.prix_achat_net and b1.remise == 0
    # Lignes problématiques signalées
    sans_marque = [lg for lg in a.lignes if any(x.champ == "marque" for x in lg.anomalies)]
    assert any("Branded sweater" in lg.nom for lg in sans_marque)
    assert not any(x.niveau == "erreur" for lg in a.lignes for x in lg.anomalies)
    # Prix cible complété depuis l'onglet « Batch 2 » pour les articles non vendus
    tibet = next(lg for lg in a.lignes if "Tibet" in lg.nom)
    assert tibet.prix_vente == D("40.00")
    # Sourcing : 9 pistes « tests » + 16 « recherche »
    assert Counter(p.origine for p in a.pistes) == {"tests": 9, "recherche": 16}


def test_import_complet_et_totaux_identiques_a_excel(
    client: TestClient, db: Session, excel_reel: bytes
) -> None:
    apercu = _apercu(client, excel_reel)
    r = client.post("/api/import/excel/commit", json={"lignes": apercu["lignes"], "pistes": apercu["pistes"]})
    assert r.status_code == 200, r.text
    assert r.json() == {"articles": 58, "batches": 5, "pistes": 25, "ventes_datees": 0}

    assert db.scalar(select(func.count(Batch.id))) == 5
    assert db.scalar(select(func.count(HistoriqueStatut.id))) == 58
    assert db.scalar(select(func.count(PisteSourcing.id))) == 25
    batch2 = db.scalar(select(Batch).where(Batch.nom == "Batch 2"))
    assert batch2 is not None and batch2.remise_globale == D("10.56")

    # Totaux de l'onglet « SYNTHÈSE GLOBALE » de l'Excel
    k = client.get("/api/compta/kpis").json()
    assert D(k["total_investi"]) == D("800.90")
    assert D(k["valeur_stock"]) == D("453.52")
    assert D(k["ca_encaisse"]) == D("814.70")
    assert D(k["marge_globale"]) == D("13.80")
    assert k["nb_total"] == 58 and k["nb_vendus"] == 26 and k["nb_en_stock"] == 32
    assert k["nb_recus"] == 49

    # Performance par batch : Batch 1 = 14 articles, 13 vendus, CA 340,40 €, marge 175,65 €
    batches = {b["nom"]: b for b in client.get("/api/batches").json()}
    b1 = batches["Batch 1"]
    assert (b1["nb_articles"], b1["nb_vendus"]) == (14, 13)
    assert D(b1["ca"]) == D("340.40") and D(b1["marge_realisee"]) == D("175.65")


def test_import_unique(client: TestClient, excel_reel: bytes) -> None:
    apercu = _apercu(client, excel_reel)
    corps = {"lignes": apercu["lignes"], "pistes": apercu["pistes"]}
    assert client.post("/api/import/excel/commit", json=corps).status_code == 200
    r = client.post("/api/import/excel/commit", json=corps)
    assert r.status_code == 422 and "qu'une fois" in r.json()["detail"]


def test_lignes_en_erreur_bloquent_la_validation(client: TestClient, db: Session, excel_reel: bytes) -> None:
    apercu = _apercu(client, excel_reel)
    apercu["lignes"][0]["prix_achat_net"] = None
    r = client.post("/api/import/excel/commit", json={"lignes": apercu["lignes"], "pistes": []})
    assert r.status_code == 422
    assert db.scalar(select(func.count(Article.id))) == 0


def test_correction_dans_apercu(client: TestClient, db: Session, excel_reel: bytes) -> None:
    apercu = _apercu(client, excel_reel)
    ligne = next(lg for lg in apercu["lignes"] if lg["nom"] == "Carhartt demi-zip sweat")
    ligne["statut"] = "en_ligne"
    ligne["marque"] = "Carhartt WIP"
    client.post("/api/import/excel/commit", json={"lignes": apercu["lignes"], "pistes": []})
    article = db.scalar(select(Article).where(Article.nom == "Carhartt demi-zip sweat"))
    assert article is not None and article.statut == "en_ligne" and article.prix_vente_reel is None


def test_fichier_invalide(client: TestClient) -> None:
    r = client.post("/api/import/excel/preview", files={"fichier": ("x.xlsx", b"pas un excel", XLSX)})
    assert r.status_code == 422


def test_export_meme_structure(client: TestClient, excel_reel: bytes) -> None:
    apercu = _apercu(client, excel_reel)
    client.post("/api/import/excel/commit", json={"lignes": apercu["lignes"], "pistes": apercu["pistes"]})
    r = client.get("/api/export/excel")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Suivi revente", "Synthèse", "Performance par batch"} <= set(wb.sheetnames)
    ws = wb["Suivi revente"]
    entetes = [c.value for c in ws[1]]
    assert entetes[:4] == ["Nom de l'article", "Taille", "Prix d'achat", "Prix de vente"]
    assert ws.max_row == 59
    # L'export se réimporte : même nombre d'articles et même CA
    relu = lire_classeur(r.content)
    assert len(relu.lignes) == 58
    assert sum(lg.prix_vente or 0 for lg in relu.lignes if lg.statut == StatutArticle.VENDU) == D("814.70")


def test_dates_de_vente_aleatoires(client: TestClient, db: Session, excel_reel: bytes) -> None:
    from zoneinfo import ZoneInfo

    apercu = _apercu(client, excel_reel)
    r = client.post(
        "/api/import/excel/commit",
        json={
            "lignes": apercu["lignes"],
            "pistes": [],
            "periode_ventes": {"du": "2026-10-01", "au": "2026-10-16"},
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["ventes_datees"] == 26
    paris = ZoneInfo("Europe/Paris")
    vendus = db.scalars(select(Article).where(Article.statut == "vendu")).all()
    for a in vendus:
        assert a.date_vente is not None
        jour = a.date_vente.astimezone(paris)
        assert "2026-10-01" <= jour.date().isoformat() <= "2026-10-16"
        assert 9 <= jour.hour < 22
        if a.date_achat:  # jamais vendu avant d'avoir été acheté
            assert jour.date() >= a.date_achat
    # Les non-vendus n'ont pas de date de vente
    assert all(a.date_vente is None for a in db.scalars(select(Article).where(Article.statut != "vendu")))
    # Toutes les ventes tombent en octobre : même CA que la synthèse Excel
    k = client.get("/api/compta/kpis", params={"du": "2026-10-01", "au": "2026-10-31"}).json()
    assert D(k["ca_encaisse"]) == D("814.70") and k["ventes_sans_date"] == 0


def test_periode_ventes_invalide(client: TestClient, excel_reel: bytes) -> None:
    apercu = _apercu(client, excel_reel)
    r = client.post(
        "/api/import/excel/commit",
        json={
            "lignes": apercu["lignes"],
            "pistes": [],
            "periode_ventes": {"du": "2026-10-16", "au": "2026-10-01"},
        },
    )
    assert r.status_code == 422
